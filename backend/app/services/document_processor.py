import os

import pandas as pd
from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader


def extract_text_from_file(file_path: str, file_type: str) -> str:
    if not os.path.exists(file_path):
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    if file_type == ".txt":
        return extract_txt(file_path)

    if file_type == ".md":
        return extract_markdown(file_path)

    if file_type == ".csv":
        return extract_csv(file_path)

    if file_type == ".pdf":
        return extract_pdf(file_path)

    if file_type == ".docx":
        return extract_docx(file_path)

    if file_type == ".pptx":
        return extract_pptx(file_path)

    raise ValueError(
        f"Unsupported file type: {file_type}"
    )


def extract_txt(file_path: str) -> str:
    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:
        return file.read()


def extract_markdown(file_path: str) -> str:
    return extract_txt(file_path)


def extract_csv(file_path: str) -> str:
    dataframe = pd.read_csv(file_path)

    return dataframe.to_csv(
        index=False
    )


def extract_pdf(file_path: str) -> str:
    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n\n".join(pages)


def extract_docx(file_path: str) -> str:
    document = DocxDocument(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    return "\n".join(paragraphs)


def extract_pptx(file_path: str) -> str:
    presentation = Presentation(file_path)

    slides = []

    for slide in presentation.slides:
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                if shape.text.strip():
                    slide_text.append(shape.text)

        if slide_text:
            slides.append("\n".join(slide_text))

    return "\n\n".join(slides)